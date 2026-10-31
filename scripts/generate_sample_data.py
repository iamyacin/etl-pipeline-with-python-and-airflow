"""Generate the deterministic, intentionally messy sample Excel exports.

The data is synthetic. A small number of rows violate the documented rules so
the quarantine path is exercised every time the pipeline runs.
"""

from __future__ import annotations

import random
from datetime import date, timedelta
from pathlib import Path

import pandas as pd
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
RAW_DIR = ROOT / "data" / "raw"
SEED = 20261031

COUNTRIES = ["US", "DE", "FR", "GB", "ES", "CA", "NL", "SE"]

PRODUCTS = [
    ("P-1001", "Wireless Headphones", "Electronics", 89.90),
    ("P-1002", "Mechanical Keyboard", "Electronics", 129.00),
    ("P-1003", "USB-C Hub", "Electronics", 49.50),
    ("P-1004", "27-inch Monitor", "Electronics", 329.00),
    ("P-1005", "Noise-cancelling Earbuds", "Electronics", 149.00),
    ("P-2001", "Espresso Grinder", "Home", 199.00),
    ("P-2002", "Cast Iron Pan", "Home", 45.00),
    ("P-2003", "Linen Bedding Set", "Home", 119.00),
    ("P-2004", "LED Desk Lamp", "Home", 39.90),
    ("P-2005", "Storage Basket", "Home", 24.00),
    ("P-3001", "Yoga Mat", "Sports", 35.00),
    ("P-3002", "Adjustable Dumbbells", "Sports", 249.00),
    ("P-3003", "Trail Running Shoes", "Sports", 139.00),
    ("P-3004", "Insulated Water Bottle", "Sports", 29.00),
    ("P-3005", "Resistance Band Set", "Sports", 22.50),
    ("P-4001", "Vitamin C Serum", "Beauty", 32.00),
    ("P-4002", "Hair Dryer", "Beauty", 79.00),
    ("P-4003", "Mineral Sunscreen", "Beauty", 27.50),
    ("P-5001", "The Pragmatic Programmer", "Books", 42.00),
    ("P-5002", "Designing Data-Intensive Applications", "Books", 54.00),
    ("P-5003", "Atomic Habits", "Books", 21.00),
    ("P-6001", "Gardening Tool Set", "Garden", 64.00),
    ("P-6002", "Ceramic Planter", "Garden", 18.50),
    ("P-6003", "Solar Garden Lights", "Garden", 34.00),
]

STATUSES = ["delivered", "shipped", "pending", "cancelled", "returned"]
STATUS_WEIGHTS = [55, 20, 10, 10, 5]
LINE_COUNTS = [1, 2, 3]
LINE_WEIGHTS = [60, 30, 10]
QUANTITIES = [1, 1, 2, 2, 3, 4]
DISCOUNTS = [0.0, 0.05, 0.10, 0.15, 0.20]
DISCOUNT_WEIGHTS = [50, 15, 20, 10, 5]

DIRTY_SPECS = [
    {"customer_id": None},
    {"quantity": 0},
    {"quantity": -2},
    {"unit_price": "N/A"},
    {"unit_price": -9.99},
    {"status": "refunded"},
    {"order_date": "not-a-date"},
    {"country": "USA"},
    {"discount": 1.5},
    {"product_id": None},
    {"order_date": "2026-13-45"},
]


def _customers(rng: random.Random) -> list[dict[str, str]]:
    return [
        {
            "customer_id": f"CUST-{index:04d}",
            "country": rng.choice(COUNTRIES),
        }
        for index in range(1, 81)
    ]


def _base_orders(rng: random.Random, customers: list[dict[str, str]], month: str) -> list[dict]:
    year, month_number = (int(part) for part in month.split("-"))
    month_start = date(year, month_number, 1)
    days_in_month = 30 if month_number in (9, 4, 6, 11) else 31
    if (year, month_number) == (2026, 10):
        days_in_month = 30  # dataset was generated on the 31st

    rows: list[dict] = []
    order_number = int(month.replace("-", "")) * 100
    order_count = 220 if month_number == 9 else 260

    for _ in range(order_count):
        order_number += 1
        order_id = f"ORD-{order_number}"
        customer = rng.choice(customers)
        order_date = month_start + timedelta(days=rng.randrange(days_in_month))
        status = rng.choices(STATUSES, weights=STATUS_WEIGHTS, k=1)[0]
        line_count = rng.choices(LINE_COUNTS, weights=LINE_WEIGHTS, k=1)[0]

        for product_id, product_name, category, unit_price in rng.sample(PRODUCTS, k=line_count):
            rows.append(
                {
                    "order_id": order_id,
                    "customer_id": customer["customer_id"],
                    "order_date": order_date,
                    "product_id": product_id,
                    "product_name": product_name,
                    "category": category,
                    "quantity": rng.choice(QUANTITIES),
                    "unit_price": unit_price,
                    "discount": rng.choices(DISCOUNTS, weights=DISCOUNT_WEIGHTS, k=1)[0],
                    "status": status,
                    "country": customer["country"],
                }
            )
    return rows


def _add_dirty_rows(
    rng: random.Random,
    rows: list[dict],
    customers: list[dict[str, str]],
    month: str,
) -> tuple[list[dict], int]:
    month_start = date.fromisoformat(f"{month}-01")
    dirty: list[dict] = []

    for index, spec in enumerate(DIRTY_SPECS, start=1):
        template = dict(rng.choice(rows))
        template.update(
            {
                "order_id": f"ORD-BAD-{month.replace('-', '')}-{index:03d}",
                "customer_id": rng.choice(customers)["customer_id"],
                "order_date": month_start + timedelta(days=index % 20),
                "product_id": f"P-BAD-{index:03d}",
                "product_name": "Invalid sample row",
                "quantity": 1,
                "unit_price": 10.0,
                "discount": 0.0,
                "status": "delivered",
                "country": "US",
            }
        )
        template.update(spec)
        dirty.append(template)

    # Exact duplicate lines exercise the deduplication rule.
    for _ in range(12):
        dirty.append(dict(rng.choice(rows)))

    combined = rows + dirty
    rng.shuffle(combined)
    return combined, len(dirty)


def _write_excel(path: Path, frame: pd.DataFrame) -> None:
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        frame.to_excel(writer, index=False, sheet_name="orders")
        worksheet = writer.sheets["orders"]
        header_fill = PatternFill("solid", fgColor="1F4E78")
        for cell in worksheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = header_fill
        for column_index, column in enumerate(frame.columns, start=1):
            longest = max(len(str(column)), *(len(str(value)) for value in frame[column].head(200)))
            width = min(longest + 2, 32)
            worksheet.column_dimensions[get_column_letter(column_index)].width = width
        worksheet.freeze_panes = "A2"


def main() -> None:
    rng = random.Random(SEED)
    customers = _customers(rng)
    RAW_DIR.mkdir(parents=True, exist_ok=True)

    for month in ("2026-09", "2026-10"):
        rows = _base_orders(rng, customers, month)
        combined, dirty_count = _add_dirty_rows(rng, rows, customers, month)
        frame = pd.DataFrame(combined)
        path = RAW_DIR / f"orders_{month.replace('-', '_')}.xlsx"
        _write_excel(path, frame)
        print(
            f"{path.relative_to(ROOT)}: {len(frame):,} rows "
            f"({dirty_count} intentionally bad/duplicate)"
        )


if __name__ == "__main__":
    main()
